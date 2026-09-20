from pathlib import Path

import pandas as pd
import json

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

from config import (
    ROOT_DIR,
    PROCESSED_FILE,
    METADATA_FILE,
    REPORT_DIR,
    FIGURE_DIR
)


def add_title(document):

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        "TWITTER SENTIMENT ANALYSIS"
    )

    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(
        38, 70, 83
    )

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = subtitle.add_run(
        "End-to-End Natural Language Processing "
        "and Machine Learning Project"
    )

    run.font.size = Pt(13)


def add_heading(document, text):

    heading = document.add_heading(
        text,
        level=1
    )

    heading.runs[0].font.color.rgb = (
        RGBColor(38, 70, 83)
    )


def add_figure(document, filename):

    path = FIGURE_DIR / filename

    if path.exists():

        document.add_picture(
            str(path),
            width=Inches(6.3)
        )

        paragraph = document.paragraphs[-1]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )


def main():

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    df = pd.read_csv(
        PROCESSED_FILE
    )

    with open(
        METADATA_FILE,
        "r",
        encoding="utf-8"
    ) as f:

        metadata = json.load(f)

    document = Document()

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    add_title(
        document
    )

    document.add_paragraph(
        "Prepared as an end-to-end Twitter/X "
        "sentiment intelligence solution."
    )

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    add_heading(
        document,
        "1. Executive Summary"
    )

    document.add_paragraph(
        f"""
This project develops an automated sentiment analysis system
for Twitter-style social media text. The system performs data
validation, text preprocessing, feature engineering, machine
learning model training, model evaluation, topic-level analysis,
visual analytics, and deployment.

The processed dataset contains {len(df):,} tweets.
The final selected model was {metadata.get('model', 'N/A')}.
"""
    )

    # --------------------------------------------------------
    # DATASET
    # --------------------------------------------------------

    add_heading(
        document,
        "2. Dataset"
    )

    document.add_paragraph(
        """
The project uses the Sentiment140 dataset. The dataset contains
Twitter messages labeled for sentiment classification. The
original collection contains approximately 1.6 million tweets.
"""
    )

    # --------------------------------------------------------
    # PREPROCESSING
    # --------------------------------------------------------

    add_heading(
        document,
        "3. Data Preprocessing"
    )

    preprocessing_points = [
        "Removed missing tweet records.",
        "Removed duplicate tweets.",
        "Decoded HTML entities.",
        "Normalized URLs.",
        "Normalized user mentions.",
        "Converted hashtags into textual terms.",
        "Normalized repeated characters.",
        "Normalized whitespace.",
        "Converted text to lowercase.",
        "Created tweet-length and word-count features."
    ]

    for item in preprocessing_points:

        document.add_paragraph(
            item,
            style="List Bullet"
        )

    # --------------------------------------------------------
    # MODEL
    # --------------------------------------------------------

    add_heading(
        document,
        "4. Machine Learning Methodology"
    )

    document.add_paragraph(
        """
The system evaluates multiple supervised machine learning
algorithms. TF-IDF word n-grams and character n-grams are used
to represent tweet text numerically. Candidate models include
Logistic Regression, Linear Support Vector Machine, and
Complement Naive Bayes.
"""
    )

    # --------------------------------------------------------
    # METRICS
    # --------------------------------------------------------

    add_heading(
        document,
        "5. Evaluation"
    )

    metrics = metadata.get(
        "metrics",
        {}
    )

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.style = "Table Grid"

    table.rows[0].cells[0].text = (
        "Metric"
    )

    table.rows[0].cells[1].text = (
        "Value"
    )

    for metric in [
        "accuracy",
        "precision",
        "recall",
        "f1"
    ]:

        row = table.add_row()

        row.cells[0].text = metric.title()

        value = metrics.get(
            metric,
            "N/A"
        )

        if isinstance(value, float):

            value = f"{value:.4f}"

        row.cells[1].text = str(
            value
        )

    # --------------------------------------------------------
    # VISUALIZATIONS
    # --------------------------------------------------------

    add_heading(
        document,
        "6. Visual Analytics"
    )

    figures = [
        "01_class_distribution.png",
        "02_tweet_length_distribution.png",
        "04_positive_wordcloud.png",
        "05_negative_wordcloud.png",
        "07_confusion_matrix.png",
        "08_roc_curve.png",
        "09_topic_sentiment.png",
        "10_topic_sentiment_heatmap.png"
    ]

    for figure in figures:

        add_figure(
            document,
            figure
        )

    # --------------------------------------------------------
    # CONCLUSION
    # --------------------------------------------------------

    add_heading(
        document,
        "7. Conclusion"
    )

    document.add_paragraph(
        """
The completed system provides an end-to-end machine learning
pipeline for classifying sentiment in Twitter-style text.
Beyond individual tweet classification, the solution supports
topic-specific sentiment analysis and professional visual
reporting.

The model should be interpreted as a statistical classifier
rather than a perfect representation of public opinion. In
particular, Sentiment140 uses automatically generated noisy
labels, which should be considered when interpreting results.
"""
    )

    # --------------------------------------------------------
    # LIMITATIONS
    # --------------------------------------------------------

    add_heading(
        document,
        "8. Limitations and Future Improvements"
    )

    limitations = [
        "Sentiment140 labels are noisy because they were generated using distant supervision.",
        "The dataset represents historical Twitter data rather than current public opinion.",
        "The system is primarily designed for English text.",
        "Sarcasm and irony remain difficult NLP problems.",
        "Topic matching currently uses keyword-based filtering.",
        "A transformer model such as RoBERTa or DeBERTa could be evaluated as a future improvement.",
        "Current X API data can be integrated for real-time monitoring where the appropriate API access is available."
    ]

    for item in limitations:

        document.add_paragraph(
            item,
            style="List Bullet"
        )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    output = REPORT_DIR / "final_report.docx"

    document.save(
        output
    )

    print(
        f"Report generated: {output}"
    )


if __name__ == "__main__":
    main()